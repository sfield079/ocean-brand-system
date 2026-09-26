"""Ocean legal documents: contracts, EPC agreements, NDAs, MSAs, signature pages (approved L5).

Plain institutional standard (brand/BRAND-SYSTEM.md §6): White field, Blackmoss only, Times New Roman
throughout, horizontal logo 24 pt on page 1, graphic mark 24 pt on continuation pages, left-aligned
body 11.5/16.5 pt, numbered sentence-case headings, one diamond as the end-of-document mark.
Reference: surfaces/formal/Ocean_Formal_Contract.pdf.

PDFs embed Times New Roman when it is installed, else the metric-identical Liberation Serif bundled in
brand/fonts/legal/. build_docx() writes an editable Word template that names Times New Roman.
"""
import json, re, os
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Flowable, KeepTogether

ROOT = Path(__file__).resolve().parent.parent
NOTICE_LABEL = {k: v['label'] for k, v in json.loads((ROOT/'brand/notices.json').read_text())['notices'].items()}
TOKENS = json.loads((ROOT/'tokens/ocean.tokens.json').read_text())
INK = colors.HexColor(TOKENS['color']['blackmoss']['value'])
WHITE = colors.HexColor(TOKENS['color']['white']['value'])

def _find(names):
    dirs = ['/usr/share/fonts', '/Library/Fonts', os.path.expanduser('~/Library/Fonts'), 'C:/Windows/Fonts',
            os.path.expanduser('~/.local/share/fonts')]
    for d in dirs:
        for root, _, files in os.walk(d) if os.path.isdir(d) else []:
            for n in names:
                if n in files: return os.path.join(root, n)
    return None
def _register():
    tnr = _find(['Times New Roman.ttf', 'times.ttf']); tnrb = _find(['Times New Roman Bold.ttf', 'timesbd.ttf'])
    reg = tnr or str(ROOT/'brand/fonts/legal/LiberationSerif-Regular.ttf')
    bold = tnrb or str(ROOT/'brand/fonts/legal/LiberationSerif-Bold.ttf')
    pdfmetrics.registerFont(TTFont('LegalRegular', reg)); pdfmetrics.registerFont(TTFont('LegalBold', bold))
    pdfmetrics.registerFontFamily('LegalRegular', normal='LegalRegular', bold='LegalBold', italic='LegalRegular', boldItalic='LegalBold')
_register()

M = 72
ST = {
  'title': ParagraphStyle('title', fontName='LegalBold', fontSize=15, leading=20, alignment=TA_CENTER, textColor=INK, spaceAfter=18),
  'body': ParagraphStyle('body', fontName='LegalRegular', fontSize=11.5, leading=16.5, textColor=INK, spaceAfter=10),
  'head': ParagraphStyle('head', fontName='LegalBold', fontSize=11.5, leading=16.5, textColor=INK, spaceBefore=8, spaceAfter=6, keepWithNext=True),
  'small': ParagraphStyle('small', fontName='LegalRegular', fontSize=8.5, leading=11, textColor=INK),
}
def P(text, s='body'): return Paragraph(escape(str(text)).replace('\n', '<br/>'), ST[s])

class Diamond(Flowable):
    def wrap(self, aw, ah): self.aw = aw; return aw, 30
    def draw(self):
        c = self.canv; x = self.aw/2; y = 15; s = 4.5; c.setFillColor(INK)
        p = c.beginPath(); p.moveTo(x, y+s); p.lineTo(x+s, y); p.lineTo(x, y-s); p.lineTo(x-s, y); p.close(); c.drawPath(p, stroke=0, fill=1)

def build(source, output):
    data = json.loads(Path(source).read_text(encoding='utf-8'))
    if not data.get('draft', False) and re.search(r'\[TBD|\[IMAGE|\[PLACEHOLDER', json.dumps(data), re.I):
        raise ValueError('Release document contains unresolved placeholders')
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output), pagesize=letter, leftMargin=M, rightMargin=M, topMargin=112, bottomMargin=72,
                            title=data['title'].replace('\n', ' '), author='Ocean RCS')
    story = [P(data['title'], 'title'), P(data['preamble'])]
    for cl in data['clauses']:
        story.append(Paragraph(f"{escape(cl['number'])}.&nbsp;&nbsp;{escape(cl['heading'])}", ST['head']))
        for num, text in cl['items']:
            t = Table([[P(num), P(text)]], colWidths=[30, letter[0]-2*M-30], hAlign='LEFT')
            t.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                                   ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.append(t)
    story += [PageBreak(), P(data['witness']), Diamond(), Spacer(1, 12)]
    w = (letter[0]-2*M-24)/2
    pa, pb = data['parties']
    cells = [[Paragraph(f"<b>{escape(pa['label'].upper())}</b>", ST['body']), '', Paragraph(f"<b>{escape(pb['label'].upper())}</b>", ST['body'])],
             [Spacer(1, 30), '', Spacer(1, 30)], [P('Signature', 'small'), '', P('Signature', 'small')],
             [P(f"Name: {pa['name']}\nTitle:\nDate:"), '', P(f"Name: {pb['name']}\nTitle:\nDate:")]]
    t = Table(cells, colWidths=[w, 24, w], hAlign='LEFT', spaceBefore=6)
    t.setStyle(TableStyle([('LINEBELOW', (0, 1), (0, 1), 0.8, INK), ('LINEBELOW', (2, 1), (2, 1), 0.8, INK),
                           ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    story.append(t)
    doc_id = f"{data['id']} · {data.get('version', 'v1.0')}"
    class Holder: total = 0
    def chrome(c, d):
        c.saveState(); w, h = letter; pg = c.getPageNumber()
        c.setFillColor(WHITE); c.rect(0, 0, w, h, stroke=0, fill=1)
        if pg == 1: c.drawImage(str(ROOT/'assets/logos/png/logo_horizontal_blackmoss.png'), M, h-48-24, width=24*510.07/142.42, height=24, mask='auto')
        else: c.drawImage(str(ROOT/'assets/logos/png/logo_graphic_blackmoss.png'), M, h-48-24, width=24, height=24, mask='auto')
        c.setFont('LegalRegular', 8); c.setFillColor(INK)
        c.drawRightString(w-M, h-58, f"Agreement No. {data['id']}")
        c.drawString(M, 40, doc_id); c.drawCentredString(w/2, 40, f"Page {pg} of {Holder.total or '—'}")
        if pg == 1:
            c.drawString(w-M-86, 40, 'Initials'); c.setStrokeColor(INK); c.setLineWidth(0.6); c.rect(w-M-54, 36, 24, 12); c.rect(w-M-24, 36, 24, 12)
        else: c.drawRightString(w-M, 40, NOTICE_LABEL.get(data.get('notice', 'confidential'), 'Confidential'))
        c.restoreState()
    # Two passes so "Page X of Y" is exact.
    import io
    tmp = SimpleDocTemplate(io.BytesIO(), pagesize=letter, leftMargin=M, rightMargin=M, topMargin=112, bottomMargin=72)
    tmp.build(list(story), onFirstPage=chrome, onLaterPages=chrome); Holder.total = tmp.page
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
    print(f'Saved {output}')

def build_docx(output):
    """Editable Word template with named styles in Times New Roman (for counsel)."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    ink = RGBColor.from_string(TOKENS['color']['blackmoss']['value'].lstrip('#'))
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ('left_margin', 'right_margin', 'bottom_margin'): setattr(sec, side, Inches(1))
    sec.top_margin = Inches(1.4); sec.different_first_page_header_footer = True
    def style(name, size, bold=False, align=None, before=0, after=8, indent=None):
        s = d.styles.add_style(name, 1); s.base_style = d.styles['Normal']
        f = s.font; f.name = 'Times New Roman'; f.size = Pt(size); f.bold = bold; f.color.rgb = ink
        pf = s.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = Pt(size*1.435)
        if align: pf.alignment = align
        if indent: pf.left_indent = Inches(indent); pf.first_line_indent = Inches(-indent)
        return s
    style('Ocean Title', 15, True, WD_ALIGN_PARAGRAPH.CENTER, after=18)
    style('Ocean Heading 1', 11.5, True, before=8, after=6)
    style('Ocean Body', 11.5)
    style('Ocean Clause 1.1', 11.5, indent=0.42)
    style('Ocean Footer', 8, after=0)
    d.styles['Normal'].font.name = 'Times New Roman'
    sec.first_page_header.paragraphs[0].add_run().add_picture(str(ROOT/'assets/logos/png/logo_horizontal_blackmoss.png'), height=Pt(24))
    sec.header.paragraphs[0].add_run().add_picture(str(ROOT/'assets/logos/png/logo_graphic_blackmoss.png'), height=Pt(24))
    for f in (sec.first_page_footer, sec.footer):
        p = f.paragraphs[0]; p.style = d.styles['Ocean Footer']; p.text = 'OCN-XXX-0000 · v1.0\tPage X of Y\tConfidential'
    d.add_paragraph('Agreement title', style='Ocean Title')
    d.add_paragraph('This Agreement is made on [date] between [Party] (“Owner”) and Ocean RCS (“Contractor”).', style='Ocean Body')
    d.add_paragraph('1.\u00a0\u00a0Definitions', style='Ocean Heading 1')
    d.add_paragraph('1.1\t“Agreement” means this agreement, including its exhibits.', style='Ocean Clause 1.1')
    d.add_paragraph('◆', style='Ocean Body').alignment = WD_ALIGN_PARAGRAPH.CENTER
    Path(output).parent.mkdir(parents=True, exist_ok=True); d.save(str(output)); print(f'Saved {output}')
