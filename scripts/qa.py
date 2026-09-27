"""PPTX checks. --draft permits placeholders, never font/geometry defects."""
import argparse, itertools, json, re
from pathlib import Path
from pptx import Presentation
from lxml import etree
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parent.parent
BRAND=json.loads((ROOT/'brand/color-system.json').read_text())
TOKENS=json.loads((ROOT/'tokens/ocean.tokens.json').read_text())
PALETTE={v['value'].lstrip('#').upper() for v in TOKENS['color'].values()}
FACES={'Stack Sans Headline','Stack Sans Headline Light','Stack Sans Headline Medium','+mn-lt','+mj-lt'}
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def check(path,draft=False):
    prs=Presentation(path); errors=[]; placeholders=0
    # Decks are built only with lib/ocean.js (AGENTS.md). Hand-drawn decks from python-pptx or other tools are rejected.
    if not (prs.core_properties.subject or '').startswith('Built with ocean-brand-system'):
        errors.append('not built with lib/ocean.js: rebuild from a content file with lib/recipe.js (AGENTS.md, Build rule)')
    W,H=prs.slide_width/914400,prs.slide_height/914400
    for i,slide in enumerate(prs.slides,1):
        boxes=[]
        for sh in slide.shapes:
            x,y,w,h=[float(v or 0)/914400 for v in (sh.left,sh.top,sh.width,sh.height)]
            if min(x,y)<-0.01 or x+w>W+0.01 or y+h>H+0.01:
                errors.append(f'slide {i}: off-slide object {sh.name}')
            if sh.has_text_frame and sh.text_frame.text.strip():
                text=sh.text_frame.text.strip()
                if x<0.55 or x+w>W-0.55 or y<0.4 or y+h>H-0.25:
                    errors.append(f'slide {i}: text violates safe area: {text[:35]}')
                boxes.append((x,y,w,h,text[:35]))
            if sh.name.startswith('ocean:logo-horizontal') and w<0.5:
                errors.append(f'slide {i}: horizontal logo below the 0.5 in minimum (BRAND-SYSTEM.md §16)')
            if sh.name.startswith('ocean:logo-graphic') and w<0.25:
                errors.append(f'slide {i}: graphic mark below the 0.25 in minimum')
            # Photographs must never be stretched: visible crop aspect must equal the frame aspect.
            if sh.shape_type==13 and (sh.name.startswith('ocean:photo') or sh.name.startswith('ocean:environment')):
                iw,ih=sh.image.size
                cl,cr,ct,cb=sh.crop_left,sh.crop_right,sh.crop_top,sh.crop_bottom
                vis=(iw*(1-cl-cr))/max(1,ih*(1-ct-cb))
                if abs(vis/(w/h)-1)>0.02: errors.append(f'slide {i}: photo distorted {vis:.2f} vs frame {w/h:.2f} ({sh.name})')
        # The close carries the logo only (BRAND-SYSTEM.md §8).
        if i==len(prs.slides) and len(prs.slides)>2:
            if any(sh.has_text_frame and sh.text_frame.text.strip() for sh in slide.shapes):
                errors.append(f'slide {i}: the close must carry the logo only, no words')
        for a,b in itertools.combinations(boxes,2):
            if min(a[0]+a[2],b[0]+b[2])-max(a[0],b[0])>0.04 and min(a[1]+a[3],b[1]+b[3])-max(a[1],b[1])>0.04:
                errors.append(f'slide {i}: text overlap: {a[4]} / {b[4]}')
    # Distribution notice on every slide but the cover and the close (decisions.md Part 10).
    NOTICES=json.loads((ROOT/'brand/notices.json').read_text())
    labels=[c['label'].split(' · ')[0].lower() for c in NOTICES['classifications'].values()]
    for i,slide in enumerate(prs.slides,1):
        if i in (1,len(prs.slides)): continue
        txt=' '.join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame).lower()
        if not any(l in txt for l in labels):
            errors.append(f'slide {i}: no distribution notice in the footer (for example Confidential · Draft · Do not use)')
    if abs(W-13.333)>0.02 or abs(H-7.5)>0.02: errors.append(f'slide size {W:.2f} x {H:.2f} in; Ocean decks are 13.333 x 7.5 in (16:9)')
    if len(prs.slides)>2:
        bg=lambda sl: sl.background.fill.fore_color.rgb if sl.background.fill.type==1 else None
        if str(bg(prs.slides[0]))!=str(bg(prs.slides[len(prs.slides)-1])):
            errors.append('cover and close use different pairings')
    with ZipFile(path) as z:
        for name in z.namelist():
            if not re.match(r'ppt/(slides/slide\d+|charts/chart\d+)\.xml$',name):continue
            xml=etree.fromstring(z.read(name))
            for text in xml.findall('.//a:t',NS):
                placeholders+=len(re.findall(r'\[TBD|\[IMAGE|\[PLACEHOLDER',text.text or '',re.I))
            for face in xml.findall('.//a:latin',NS):
                if face.get('typeface') not in FACES:
                    errors.append(f'{name}: unapproved font {face.get("typeface")}')
            for run in xml.findall('.//a:rPr',NS):
                if run.get('sz') and int(run.get('sz'))<900: errors.append(f'{name}: text below 9 pt')
                rgb=run.find('./a:solidFill/a:srgbClr',NS)
                if rgb is not None and rgb.get('val').upper() not in PALETTE:
                    errors.append(f'{name}: off-palette text {rgb.get("val")}')
            # Verify the theme cannot silently introduce an unapproved default.
        theme=etree.fromstring(z.read('ppt/theme/theme1.xml'))
        for face in theme.findall('.//a:fontScheme/a:majorFont/a:latin',NS)+theme.findall('.//a:fontScheme/a:minorFont/a:latin',NS):
            if face.get('typeface')!='Stack Sans Headline': errors.append('Theme font is not Stack Sans Headline')
    if placeholders and not draft: errors.append(f'{placeholders} unresolved placeholders in release')
    print(f'QA: {path} ({len(prs.slides)} slides); placeholders: {placeholders}')
    for error in errors:print('ERROR:',error)
    print('RESULT FAIL' if errors else 'RESULT PASS (automated checks only; visual review still required)')
    return not errors
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('path');p.add_argument('--draft',action='store_true');a=p.parse_args()
    raise SystemExit(0 if check(a.path,a.draft) else 1)
