import json, tempfile, unittest, sys
from pathlib import Path
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
from lib.publication import build
from scripts.qa_pdf import check

class PublicationTests(unittest.TestCase):
    def test_embedded_font_and_live_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'report.pdf'
            build(ROOT/'documents/_starter/report.json',out)
            self.assertTrue(check(out,draft=True))
    def test_release_blocks_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/'source.json'
            src.write_text(json.dumps({'title':'[TBD]','sections':[]}))
            with self.assertRaisesRegex(ValueError,'placeholder'):build(src,Path(tmp)/'bad.pdf')


class LegalTests(unittest.TestCase):
    def test_legal_pdf_uses_times_metric_font(self):
        from lib import legal
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'contract.pdf'
            legal.build(ROOT/'documents/_starter/contract.json',out)
            self.assertTrue(check(out,draft=True,legal=True))
            self.assertFalse(check(out,draft=True))  # Stack Sans-only check must reject legal fonts

class DeckQATests(unittest.TestCase):
    def test_stretched_photo_is_rejected(self):
        import subprocess, shutil
        from pptx import Presentation
        from pptx.util import Inches
        sys.path.insert(0, str(ROOT/'scripts'))
        from scripts import qa
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
            pic = s.shapes.add_picture(str(ROOT/'assets/images/rooftop.jpg'), Inches(1), Inches(1), Inches(4), Inches(4))
            pic.name = 'ocean:photo'
            out = Path(tmp)/'t.pptx'; prs.save(out)
            self.assertFalse(qa.check(out, draft=True))

if __name__=='__main__':unittest.main()

class LegalTests(unittest.TestCase):
    def test_legal_pdf_uses_times_metric_font(self):
        from lib import legal
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'contract.pdf'
            legal.build(ROOT/'documents/_starter/contract.json',out)
            self.assertTrue(check(out,draft=True,legal=True))
            self.assertFalse(check(out,draft=True))  # Stack Sans-only check must reject legal fonts
